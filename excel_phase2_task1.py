import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime

xlsm_path = r"C:\Users\LENOVO\Desktop\Excel Project\excel project 1.xlsm"

# Load workbook
wb = openpyxl.load_workbook(xlsm_path)

print("=== TASK 1: VERIFY INSIGHTS SHEET ===\n")

# Check if Insights sheet exists
if "Insights" in wb.sheetnames:
    ws = wb["Insights"]
    print("✓ Insights sheet EXISTS")
    print(f"  Current content dimensions: {ws.dimensions}")
    
    # Check first few cells
    print(f"  A1: {ws['A1'].value}")
    print(f"  A2: {ws['A2'].value}")
    print(f"  A3: {ws['A3'].value}")
else:
    print("✗ Insights sheet DOES NOT EXIST - needs to be created")

print("\n=== TASK 2: AUDIT DASHBOARD ===\n")
dashboard = wb["Dashboard"]
print(f"Dashboard dimensions: {dashboard.dimensions}")

# Check for some key cells that might have formulas
print("\nSample cells in Dashboard (A1:F10):")
for row in range(1, 11):
    for col in range(1, 7):
        cell = dashboard.cell(row, col)
        if cell.value:
            print(f"  {cell.coordinate}: {cell.value[:60] if isinstance(cell.value, str) else cell.value}")

print("\n=== DATA VERIFICATION ===\n")
data_ws = wb["Data"]
print(f"Data sheet: {data_ws.dimensions}")
print(f"Data rows (excluding header): {data_ws.max_row - 1}")

# Sample data verification
print("\nFirst data row values (row 2):")
for col in range(1, 14):
    val = data_ws.cell(2, col).value
    print(f"  Col {col}: {val}")

# Check Amount column (F) - should have numeric values
print(f"\nAmount column (F) - sample values from rows 2-5:")
for row in range(2, 6):
    val = data_ws.cell(row, 6).value
    print(f"  Row {row}: {val}")

