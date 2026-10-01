import openpyxl
from openpyxl.utils import get_column_letter
from pathlib import Path
import os

xlsm_path = r"C:\Users\LENOVO\Desktop\Excel Project\excel project 1.xlsm"

# Load workbook
wb = openpyxl.load_workbook(xlsm_path)

print("=== CURRENT WORKBOOK STRUCTURE ===")
print("Sheets:", list(wb.sheetnames))

# Inspect Data sheet
data_ws = wb["Data"]
print(f"\nData Sheet: {data_ws.dimensions}")
print(f"Max Row: {data_ws.max_row}, Max Col: {data_ws.max_column}")

# Get headers
print("\nHeaders (first 13 columns):")
for col in range(1, 14):
    cell_value = data_ws.cell(1, col).value
    print(f"  Col {col}: {cell_value}")

# Check for tables
print(f"\nTables in Data sheet: {[t.name for t in data_ws.tables.values()]}")

# Check what sheets exist
print("\n=== SHEET ANALYSIS ===")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"{sheet_name}: {ws.dimensions}")

print("\n✓ Analysis complete")
